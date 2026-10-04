#!/usr/bin/env python3
"""Herramienta para construir decks VivaSlide a partir de la plantilla oficial.

Subcomandos:
  inspect  <deck.pptx>                    Lista cada slide con sus shapes (nombre, posición, texto).
  build    <spec.json> -o <salida.pptx>   Arma un deck nuevo desde la plantilla según un spec JSON.

El spec JSON tiene esta forma:

{
  "template": "assets/vivaslide-template.pptx",   (opcional; por defecto la plantilla oficial)
  "footer": "VIVA CONSULTING EMPRESAS  ·  ANÁLISIS CORPORATIVO · <CLIENTE> · <TEMA> · DIRECTORIO  ·  CONFIDENCIAL",
  "slides": [
    {"from": 1, "text": {"Text 3": "Cliente — Tema"}},
    {"from": 5, "text": {"Text 3": "Título-hallazgo"},
     "charts": {"Chart 0": {"categories": ["A", "B"], "series": {"2025": [1, 2], "2026": [3, 4]}}}},
    {"from": 6, "tables": {"Table 0": [["Zona", "2025"], ["A", "1"]]}},
    {"from": 13, "text": {"Text 12": ["Párrafo 1", "Párrafo 2"]}},
    {"from": 14}
  ]
}

- "from" = número (1-based) de la slide de la plantilla que se clona.
- "text": shape -> string (un párrafo) o lista de strings (un párrafo por elemento).
  Se conserva el formato del primer run de cada párrafo de la plantilla.
- "charts": reemplaza los datos del gráfico nativo (sigue siendo editable en PowerPoint).
- "tables": matriz de strings; se escribe sobre la tabla existente (filas/columnas
  sobrantes se eliminan; si faltan filas se clonan de la penúltima).
- "delete": lista de nombres de shapes a eliminar (p. ej. una tarjeta que sobra).
- El pie "Text N" con formato "k/N" se renumera solo y el texto largo de pie se reemplaza por "footer".
"""
import argparse
import copy
import json
import os
import re
import sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_TEMPLATE = os.path.join(HERE, "..", "assets", "vivaslide-template.pptx")
PAGE_RE = re.compile(r"^\s*\d+\s*/\s*\d+\s*$")
FOOTER_PREFIX = "VIVA CONSULTING EMPRESAS"


def inspect(path):
    prs = Presentation(path)
    for i, slide in enumerate(prs.slides, 1):
        print(f"--- slide {i}")
        for sh in slide.shapes:
            txt = ""
            if sh.has_text_frame:
                txt = sh.text_frame.text.replace("\n", " | ")[:70]
            kind = "CHART" if sh.has_chart else "TABLE" if sh.has_table else str(sh.shape_type).split(" ")[0]
            print(f"  {sh.name:10} {kind:12} x={sh.left/914400:.2f} y={sh.top/914400:.2f} "
                  f"w={sh.width/914400:.2f} h={sh.height/914400:.2f}  {txt}")


def set_paragraphs(text_frame, values):
    """Escribe values (str o lista) conservando el formato de los párrafos de la plantilla."""
    if isinstance(values, str):
        values = [values]
    paras = list(text_frame.paragraphs)
    txBody = text_frame._txBody
    # Asegura tantos párrafos como valores, clonando el último párrafo modelo.
    while len(paras) < len(values):
        new_p = copy.deepcopy(paras[-1]._p)
        txBody.append(new_p)
        paras = list(text_frame.paragraphs)
    for p, val in zip(paras, values):
        runs = p.runs
        if runs:
            runs[0].text = val
            for r in runs[1:]:
                r._r.getparent().remove(r._r)
        else:
            p.add_run().text = val
    for p in paras[len(values):]:
        txBody.remove(p._p)


def fill_table(table, rows):
    tbl = table._tbl
    tr_list = tbl.tr_lst
    # Ajusta número de filas: clona la penúltima fila (fila de datos típica) si faltan.
    while len(tbl.tr_lst) < len(rows):
        model = tbl.tr_lst[-2] if len(tbl.tr_lst) > 1 else tbl.tr_lst[-1]
        model.addnext(copy.deepcopy(model))
    while len(tbl.tr_lst) > len(rows):
        # Elimina filas intermedias para conservar el formato de la última (TOTAL).
        victim = tbl.tr_lst[-2] if len(tbl.tr_lst) > 2 else tbl.tr_lst[-1]
        tbl.remove(victim)
    ncols = max(len(r) for r in rows)
    for r_i, row in enumerate(rows):
        for c_i in range(len(table.columns)):
            cell = table.cell(r_i, c_i)
            val = row[c_i] if c_i < len(row) else ""
            set_paragraphs(cell.text_frame, str(val))
    # Elimina columnas sobrantes
    while len(table.columns) > ncols:
        grid = tbl.tblGrid
        grid.remove(grid.gridCol_lst[-1])
        for tr in tbl.tr_lst:
            tr.remove(tr.tc_lst[-1])


def fill_chart(chart, spec):
    data = CategoryChartData()
    data.categories = spec["categories"]
    for name, values in spec["series"].items():
        data.add_series(name, values)
    chart.replace_data(data)


def duplicate_slide(prs, src_index):
    """Clona una slide (con sus gráficos, que se copian como partes nuevas)."""
    src = prs.slides[src_index]
    dst = prs.slides.add_slide(src.slide_layout)
    for shp in list(dst.shapes):
        shp._element.getparent().remove(shp._element)
    rel_map = {}
    for rel in src.part.rels.values():
        if "notesSlide" in rel.reltype:
            continue
        target = rel._target
        if "chart" in rel.reltype:
            target = clone_chart_part(prs, target)
        if rel.is_external:
            new_rid = dst.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            new_rid = dst.part.rels.get_or_add(rel.reltype, target)
        rel_map[rel.rId] = new_rid
    for el in src.shapes._spTree.iterchildren():
        tag = el.tag.split("}")[1]
        if tag in ("nvGrpSpPr", "grpSpPr"):
            continue
        new_el = copy.deepcopy(el)
        for node in new_el.iter():
            for attr in list(node.attrib):
                if attr.endswith("}embed") or attr.endswith("}id") or attr.endswith("}link"):
                    old = node.attrib[attr]
                    if old in rel_map:
                        node.attrib[attr] = rel_map[old]
        dst.shapes._spTree.append(new_el)
    # Fondo de la slide
    src_bg = src._element.cSld.bg
    if src_bg is not None:
        dst._element.cSld.insert(0, copy.deepcopy(src_bg))
    return dst


_chart_counter = [0]


def clone_chart_part(prs, chart_part):
    from pptx.opc.packuri import PackURI
    from pptx.parts.chart import ChartPart
    _chart_counter[0] += 1
    existing = {str(p.partname) for p in prs.part.package.iter_parts()}
    n = 100 + _chart_counter[0]
    while f"/ppt/charts/chart{n}.xml" in existing:
        n += 1
    partname = PackURI(f"/ppt/charts/chart{n}.xml")
    new_part = ChartPart.load(partname, chart_part.content_type, prs.part.package, chart_part.blob)
    # copia el Excel embebido para que "Editar datos" funcione
    for rel in chart_part.rels.values():
        if rel.is_external:
            new_part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            tgt = rel._target
            if "package" in rel.reltype:
                from pptx.opc.package import Part
                xlsx_name = PackURI(f"/ppt/embeddings/vivaslide_chart{n}.xlsx")
                tgt = Part(xlsx_name, tgt.content_type, prs.part.package, tgt.blob)
            new_rid = new_part.rels.get_or_add(rel.reltype, tgt)
            if new_rid != rel.rId:
                xml = new_part._element
                for node in xml.iter():
                    for attr in list(node.attrib):
                        if attr.endswith("}id") and node.attrib[attr] == rel.rId:
                            node.attrib[attr] = new_rid
    return new_part


def build(spec_path, out):
    spec = json.load(open(spec_path, encoding="utf-8"))
    base = os.path.dirname(os.path.abspath(spec_path))
    tpl = spec.get("template") or DEFAULT_TEMPLATE
    if not os.path.isabs(tpl) and not os.path.exists(tpl):
        tpl = os.path.join(base, tpl)
    prs = Presentation(tpl)
    n_template = len(prs.slides)

    new_slides = []
    for s in spec["slides"]:
        src = int(s["from"]) - 1
        if not 0 <= src < n_template:
            sys.exit(f"'from' fuera de rango: {s['from']} (plantilla tiene {n_template})")
        new_slides.append((duplicate_slide(prs, src), s))

    # Elimina las slides originales de la plantilla
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst)[:n_template]:
        prs.part.drop_rel(sldId.rId)
        sldIdLst.remove(sldId)

    total_numbered = sum(
        1 for sl, _ in new_slides
        if any(sh.has_text_frame and PAGE_RE.match(sh.text_frame.text) for sh in sl.shapes)
    )
    k = 0
    footer = spec.get("footer")
    for slide, s in new_slides:
        shapes = {sh.name: sh for sh in slide.shapes}
        has_page = False
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            t = sh.text_frame.text
            if PAGE_RE.match(t):
                has_page = True
                page_shape = sh
            elif footer and t.startswith(FOOTER_PREFIX):
                set_paragraphs(sh.text_frame, footer)
        if has_page:
            k += 1
            set_paragraphs(page_shape.text_frame, f"{k}/{total_numbered}")
        for name, val in s.get("text", {}).items():
            if name not in shapes:
                sys.exit(f"slide from={s['from']}: no existe el shape '{name}'. Usa 'inspect'.")
            set_paragraphs(shapes[name].text_frame, val)
        for name, rows in s.get("tables", {}).items():
            fill_table(shapes[name].table, rows)
        for name, cspec in s.get("charts", {}).items():
            fill_chart(shapes[name].chart, cspec)
        for name in s.get("delete", []):
            el = shapes[name]._element
            el.getparent().remove(el)
        if "notes" in s:
            slide.notes_slide.notes_text_frame.text = s["notes"]

    prs.save(out)
    print(f"OK -> {out} ({len(new_slides)} slides)")


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    a = sub.add_parser("inspect")
    a.add_argument("deck", nargs="?", default=DEFAULT_TEMPLATE)
    b = sub.add_parser("build")
    b.add_argument("spec")
    b.add_argument("-o", "--out", required=True)
    args = ap.parse_args()
    if args.cmd == "inspect":
        inspect(args.deck)
    else:
        build(args.spec, args.out)


if __name__ == "__main__":
    main()
